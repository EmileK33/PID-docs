"""S2-K — Async Export Worker acceptance tests (US-017).

Run with:  pytest tests/integration/test_export_worker.py -v

These tests prove the Celery export task in full isolation — no live DB, S3, or
Redis, and no sibling Phase-2 session (S2-D etc.) needs to be merged. Every
external dependency is mocked:

* **DB** — ``SessionLocal`` is patched to return a hand-rolled ``FakeSession``
  that records ``with_for_update`` / ``add`` / ``flush`` / ``commit`` / ``close``
  and dispatches ``query(Model)`` by class. No SQLite/Postgres is touched
  (the §1.2 models use Postgres-only types: ``UUID``, ``JSONB``).
* **S3** — ``put_object`` is patched and its bytes captured for CSV/XLSX
  assertions.
* **Redis** — ``_redis_client`` is patched (autouse) so notifications never hit
  the network; tests that assert the payload patch ``publish_export_event``.

The module self-configures ``sys.path`` so ``app`` (rooted at ``backend/``) is
importable and sets every §1.12 "Refuse to start" env var *before* importing
``app`` (deferred into an autouse fixture) so ``app.config`` does not
``sys.exit(1)`` and so app.* stays out of sys.modules during collection (keeps
S0-B's smoke suite green).
"""
from __future__ import annotations

import ast
import io
import os
import sys
import uuid
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from celery.exceptions import Retry

# --- make `app` (backend/app) importable -----------------------------------
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
BACKEND = REPO_ROOT / "backend"
for _p in (str(BACKEND), str(REPO_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

# --- §1.12 env contract: set BEFORE importing app code ----------------------
_REQUIRED_ENV = {
    "DATABASE_URL": "postgresql://u:p@localhost:5432/pid",
    "REDIS_URL": "redis://localhost:6379/0",
    "S3_BUCKET_NAME": "pidtest-bucket",
    "S3_REGION": "us-east-1",
    "SUPABASE_URL": "https://example.supabase.co",
    "SUPABASE_SERVICE_ROLE_KEY": "service-role-key",
    "JWT_RS256_PUBLIC_KEY": "-----BEGIN PUBLIC KEY-----\nMIIB\n-----END PUBLIC KEY-----",
    "STRIPE_SECRET_KEY": "sk_test_dummy",
    "STRIPE_WEBHOOK_SECRET": "whsec_dummy",
    "SENDGRID_API_KEY": "SG.dummy",
    "HMAC_SERVER_SECRET": "high-entropy-secret",
    "ML_MODEL_S3_KEY": "models/pid/v1.onnx",
    "ODA_CONVERTER_PATH": "/opt/oda/ODAFileConverter",
}
for _k, _v in _REQUIRED_ENV.items():
    os.environ.setdefault(_k, _v)

# Deferred app imports — bound at test-run time by the autouse fixture so app.*
# stays out of sys.modules during collection.
tasks = None  # type: ignore[assignment]
settings = None  # type: ignore[assignment]
ExportRecord = None  # type: ignore[assignment]
StoredFile = None  # type: ignore[assignment]
DetectedSymbol = None  # type: ignore[assignment]
celery_app = None  # type: ignore[assignment]

BUCKET = _REQUIRED_ENV["S3_BUCKET_NAME"]

EXPECTED_COLUMNS = [
    "id",
    "drawing_id",
    "page_number",
    "entity_class_id",
    "subtype",
    "tag_label",
    "confidence",
    "bbox_x",
    "bbox_y",
    "bbox_w",
    "bbox_h",
    "source",
]

EXPORT_ID = "aaaaaaaa-0000-0000-0000-000000000001"
DRAWING_ID = "bbbbbbbb-0000-0000-0000-000000000001"
USER_ID = "cccccccc-0000-0000-0000-000000000001"
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


# ---------------------------------------------------------------------------
# Deferred import + autouse Redis isolation
# ---------------------------------------------------------------------------
@pytest.fixture(autouse=True)
def _deferred_app_import():
    """Bind the app exports onto module globals at test-run time."""
    from app.config import settings as _settings
    from app.db.models import DetectedSymbol as _DetectedSymbol
    from app.db.models import ExportRecord as _ExportRecord
    from app.db.models import StoredFile as _StoredFile
    from app.workers.celery_app import celery_app as _celery_app
    from app.workers.export import tasks as _tasks

    g = globals()
    g["tasks"] = _tasks
    g["settings"] = _settings
    g["ExportRecord"] = _ExportRecord
    g["StoredFile"] = _StoredFile
    g["DetectedSymbol"] = _DetectedSymbol
    g["celery_app"] = _celery_app
    yield


@pytest.fixture(autouse=True)
def _isolate_redis(monkeypatch, _deferred_app_import):
    """Never touch a real Redis: replace the sync client factory with a mock.

    Tests that assert the published payload patch ``publish_export_event``
    directly; this guard protects the rest from network access.
    """
    monkeypatch.setattr(tasks, "_redis_client", lambda: MagicMock(name="redis"))
    yield


# ---------------------------------------------------------------------------
# Test doubles
# ---------------------------------------------------------------------------
class _QueryStub:
    """Minimal SQLAlchemy ``Query`` stand-in dispatching by model class."""

    def __init__(self, session: "FakeSession", model):
        self._session = session
        self._model = model

    def filter(self, *clauses, **_kw):
        self._session.filters.setdefault(self._model, []).extend(
            str(c) for c in clauses
        )
        return self

    def order_by(self, *clauses, **_kw):
        self._session.order_bys.setdefault(self._model, []).extend(
            str(c) for c in clauses
        )
        return self

    def with_for_update(self, *_a, **_kw):
        self._session.with_for_update_models.append(self._model)
        return self

    def one_or_none(self):
        return self._session._one_for(self._model)

    def first(self):
        return self._session._one_for(self._model)

    def all(self):
        return self._session._all_for(self._model)


class FakeSession:
    """Records the persistence operations the worker performs."""

    def __init__(self, record=None, symbols=None, existing_stored_file=None):
        self.record = record
        self.symbols = list(symbols or [])
        self.existing_stored_file = existing_stored_file
        self.added: list = []
        self.commit_count = 0
        self.statuses_at_commit: list = []
        self.closed = False
        self.with_for_update_models: list = []
        self.filters: dict = {}
        self.order_bys: dict = {}

    # query dispatch -------------------------------------------------------
    def query(self, model):
        return _QueryStub(self, model)

    def _one_for(self, model):
        if model is ExportRecord:
            return self.record
        if model is StoredFile:
            return self.existing_stored_file
        return None

    def _all_for(self, model):
        if model is DetectedSymbol:
            return list(self.symbols)
        return []

    # unit-of-work ---------------------------------------------------------
    def add(self, obj):
        self.added.append(obj)

    def flush(self):
        # The DB assigns ids via server_default; emulate that for new rows.
        for obj in self.added:
            if getattr(obj, "id", None) is None:
                obj.id = uuid.uuid4()

    def commit(self):
        self.commit_count += 1
        if self.record is not None:
            self.statuses_at_commit.append(self.record.status)

    def close(self):
        self.closed = True

    @property
    def added_stored_files(self):
        return [o for o in self.added if isinstance(o, StoredFile)]


def _make_record(status="Queued", **overrides):
    data = dict(
        id=EXPORT_ID,
        drawing_id=DRAWING_ID,
        user_id=USER_ID,
        format="csv",
        status=status,
        stored_file_id=None,
        initiated_at=datetime(2024, 1, 1, 12, 0, 0, tzinfo=timezone.utc),
        completed_at=None,
    )
    data.update(overrides)
    return SimpleNamespace(**data)


def _make_symbols(count_non_rejected=5):
    """Return ``count_non_rejected`` non-rejected symbols, pre-ordered as the
    worker's ``page_number ASC, id ASC`` query would return them.

    The two rejected fixtures from the brief are intentionally *not* returned
    here — the ``rejected = FALSE`` filter lives in the SQL the worker issues,
    which the FakeSession verifies separately."""
    symbols = []
    for i in range(count_non_rejected):
        symbols.append(
            SimpleNamespace(
                id=f"dddddddd-0000-0000-0000-00000000000{i}",
                drawing_id=DRAWING_ID,
                entity_class_id="valve_gate",
                subtype="gate",
                tag_label=f"V-{i}",
                confidence=0.95,
                bbox={"x": 10.0, "y": 20.0, "w": 30.0, "h": 40.0},
                source="ml",
                rejected=False,
                page_number=1,
            )
        )
    return symbols


def _make_self(retries=0):
    """A stand-in for the bound Celery task ``self``."""
    return SimpleNamespace(
        request=SimpleNamespace(retries=retries),
        retry=MagicMock(name="retry", side_effect=Retry()),
    )


@pytest.fixture()
def patch_session(monkeypatch):
    """Patch ``SessionLocal`` to return a single shared FakeSession."""

    def _install(session: FakeSession):
        monkeypatch.setattr(tasks, "SessionLocal", lambda: session)
        return session

    return _install


@pytest.fixture()
def captured_put(monkeypatch):
    """Patch ``put_object`` and capture (bucket, key, data, content_type)."""
    calls: list = []

    def _put(bucket, key, data, content_type, **_kw):
        calls.append(
            SimpleNamespace(
                bucket=bucket, key=key, data=data, content_type=content_type
            )
        )

    mock = MagicMock(side_effect=_put)
    monkeypatch.setattr(tasks, "put_object", mock)
    return SimpleNamespace(mock=mock, calls=calls)


# ===========================================================================
# AC-1 — state machine / idempotency
# ===========================================================================
def test_task_transitions_to_generating_on_pickup(patch_session, captured_put):
    record = _make_record(status="Queued")
    session = patch_session(FakeSession(record=record, symbols=_make_symbols()))

    # Capture the record status at the moment generation (S3 upload) runs, to
    # prove the Generating transition precedes any file generation.
    status_at_upload = {}

    def _put(bucket, key, data, content_type, **_kw):
        status_at_upload["status"] = record.status

    captured_put.mock.side_effect = _put

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    assert status_at_upload["status"] == "Generating"
    # First commit persisted the Generating transition before generation.
    assert session.statuses_at_commit[0] == "Generating"


def test_task_is_idempotent_when_already_generating(patch_session, captured_put):
    record = _make_record(status="Generating")
    session = patch_session(FakeSession(record=record))

    # Fresh delivery (retries=0 → allow_generating=False): a Generating row is
    # a duplicate delivery and must be skipped without re-processing.
    tasks._run_export(_make_self(retries=0), EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    captured_put.mock.assert_not_called()
    assert record.status == "Generating"
    assert session.commit_count == 0  # never transitioned / regenerated


def test_task_is_idempotent_when_already_complete(patch_session, captured_put):
    record = _make_record(status="Complete")
    session = patch_session(FakeSession(record=record))

    tasks._run_export(_make_self(retries=0), EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    captured_put.mock.assert_not_called()
    assert record.status == "Complete"
    assert session.commit_count == 0


# ===========================================================================
# AC-2 — CSV generation
# ===========================================================================
def test_csv_export_header_columns():
    data = tasks.generate_csv_bytes(_make_symbols(0))
    text = data.decode("utf-8")
    header = text.splitlines()[0]
    assert header.split(",") == EXPECTED_COLUMNS


def test_csv_export_data_rows_exclude_rejected_ordered(patch_session, captured_put):
    """AC-2 ordering + AC-9 rejected exclusion.

    The FakeSession returns exactly the 5 non-rejected symbols its
    ``page_number ASC, id ASC`` query would return; the test also asserts the
    worker issued that query with the ``rejected`` filter and the correct
    order — i.e. exclusion is requested of the DB, not done post-hoc.
    """
    symbols = _make_symbols(5)
    record = _make_record(status="Queued")
    session = patch_session(FakeSession(record=record, symbols=symbols))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    # The query the worker built for detected_symbol.
    sym_filters = " ".join(session.filters.get(DetectedSymbol, []))
    assert "rejected" in sym_filters.lower()
    sym_order = session.order_bys.get(DetectedSymbol, [])
    assert "page_number" in sym_order[0].lower()
    assert sym_order[1].lower().endswith("id asc") or "id" in sym_order[1].lower()

    # The CSV body: header + exactly 5 data rows in fixture order.
    data = captured_put.calls[0].data.decode("utf-8")
    rows = [r for r in data.splitlines() if r]
    assert len(rows) == 1 + 5
    assert rows[1].split(",")[0] == symbols[0].id
    assert rows[-1].split(",")[0] == symbols[-1].id


def test_csv_s3_upload_key(patch_session, captured_put):
    record = _make_record(status="Queued")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    call = captured_put.calls[0]
    assert call.key == f"exports/{EXPORT_ID}.csv"
    assert call.bucket == BUCKET
    assert call.content_type == "text/csv"


def test_stored_file_record_created_csv(patch_session, captured_put):
    record = _make_record(status="Queued")
    session = patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    stored = session.added_stored_files
    assert len(stored) == 1
    sf = stored[0]
    assert sf.bucket == BUCKET
    assert sf.object_key == f"exports/{EXPORT_ID}.csv"
    assert sf.file_type == "text/csv"
    assert len(sf.sha256_hash) == 64  # hex sha256
    assert sf.size_bytes == len(captured_put.calls[0].data)


def test_export_record_set_complete_atomically(patch_session, captured_put):
    record = _make_record(status="Queued")
    session = patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    assert record.status == "Complete"
    assert record.stored_file_id is not None
    assert record.completed_at is not None
    # The StoredFile insert and the Complete update commit together (final
    # commit) — Generating first, then Complete.
    assert session.statuses_at_commit == ["Generating", "Complete"]
    assert record.stored_file_id == session.added_stored_files[0].id


# ===========================================================================
# AC-3 — XLSX generation
# ===========================================================================
def test_xlsx_export_symbols_worksheet_columns(patch_session, captured_put):
    from openpyxl import load_workbook

    record = _make_record(status="Queued", format="xlsx")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "xlsx", allow_generating=False)

    wb = load_workbook(io.BytesIO(captured_put.calls[0].data))
    assert "Symbols" in wb.sheetnames
    ws = wb["Symbols"]
    header = [c.value for c in ws[1]]
    assert header == EXPECTED_COLUMNS


def test_xlsx_s3_upload_key(patch_session, captured_put):
    record = _make_record(status="Queued", format="xlsx")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "xlsx", allow_generating=False)

    call = captured_put.calls[0]
    assert call.key == f"exports/{EXPORT_ID}.xlsx"
    assert call.bucket == BUCKET
    assert call.content_type == XLSX_CONTENT_TYPE


def test_stored_file_record_created_xlsx(patch_session, captured_put):
    record = _make_record(status="Queued", format="xlsx")
    session = patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "xlsx", allow_generating=False)

    sf = session.added_stored_files[0]
    assert sf.file_type == XLSX_CONTENT_TYPE
    assert sf.object_key == f"exports/{EXPORT_ID}.xlsx"


def test_xlsx_export_data_rows_exclude_rejected(patch_session, captured_put):
    from openpyxl import load_workbook

    symbols = _make_symbols(5)
    record = _make_record(status="Queued", format="xlsx")
    patch_session(FakeSession(record=record, symbols=symbols))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "xlsx", allow_generating=False)

    ws = load_workbook(io.BytesIO(captured_put.calls[0].data))["Symbols"]
    # Header + exactly 5 data rows (rejected excluded by the SQL filter).
    assert ws.max_row == 1 + 5
    assert ws.cell(row=2, column=1).value == symbols[0].id


def test_xlsx_p1_stub_comment_present():
    src = Path(tasks.__file__).read_text(encoding="utf-8")
    assert "TODO P1 (US-022)" in src
    assert "Tables" in src  # the deferred second worksheet


# ===========================================================================
# AC-4 — Redis completion notification
# ===========================================================================
def test_redis_notification_published_on_success(patch_session, captured_put, monkeypatch):
    record = _make_record(status="Queued")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))
    publish = MagicMock(name="publish_export_event")
    monkeypatch.setattr(tasks, "publish_export_event", publish)

    tasks._run_export(_make_self(retries=0), EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    publish.assert_called_once()
    channel, payload = publish.call_args.args
    assert channel == f"export:status:{EXPORT_ID}"
    assert payload["event"] == "export_complete"
    assert payload["export_id"] == EXPORT_ID
    assert payload["drawing_id"] == DRAWING_ID
    assert payload["user_id"] == USER_ID
    assert payload["format"] == "csv"
    assert payload["timestamp"].endswith("Z")


# ===========================================================================
# AC-5 / AC-6 — retry + failure
# ===========================================================================
def test_celery_retries_on_first_exception(patch_session, captured_put, monkeypatch):
    record = _make_record(status="Queued")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))
    captured_put.mock.side_effect = RuntimeError("S3 down")
    publish = MagicMock(name="publish_export_event")
    monkeypatch.setattr(tasks, "publish_export_event", publish)

    task_self = _make_self(retries=0)
    with pytest.raises(Retry):
        tasks._run_export(task_self, EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    task_self.retry.assert_called_once()
    assert task_self.retry.call_args.kwargs["countdown"] == 30
    assert "exc" in task_self.retry.call_args.kwargs
    # AC-5: NOT marked Failed during the retry window — stays Generating.
    assert record.status == "Generating"
    publish.assert_not_called()


def test_export_record_set_failed_after_max_retries(patch_session, captured_put, monkeypatch):
    record = _make_record(status="Generating")  # set by attempt 1
    patch_session(FakeSession(record=record, symbols=_make_symbols()))
    captured_put.mock.side_effect = RuntimeError("S3 down")
    monkeypatch.setattr(tasks, "publish_export_event", MagicMock())

    task_self = _make_self(retries=1)  # retries exhausted
    with pytest.raises(RuntimeError):
        tasks._run_export(task_self, EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    task_self.retry.assert_not_called()
    assert record.status == "Failed"
    assert record.completed_at is None
    assert record.stored_file_id is None


def test_redis_failure_notification_published(patch_session, captured_put, monkeypatch):
    record = _make_record(status="Generating")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))
    captured_put.mock.side_effect = RuntimeError("S3 down")
    publish = MagicMock(name="publish_export_event")
    monkeypatch.setattr(tasks, "publish_export_event", publish)

    with pytest.raises(RuntimeError):
        tasks._run_export(_make_self(retries=1), EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    publish.assert_called_once()
    channel, payload = publish.call_args.args
    assert channel == f"export:status:{EXPORT_ID}"
    assert payload["event"] == "export_failed"
    assert payload["export_id"] == EXPORT_ID
    assert payload["drawing_id"] == DRAWING_ID
    assert payload["user_id"] == USER_ID
    assert payload["format"] == "csv"


# ===========================================================================
# AC-7 — no analytics fired from the worker
# ===========================================================================
def test_no_analytics_events_fired_from_worker():
    """The worker must not import or call any analytics emitter (§1.10)."""
    src = Path(tasks.__file__).read_text(encoding="utf-8")
    tree = ast.parse(src)
    imported_modules: list[str] = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported_modules.extend(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported_modules.append(node.module)

    joined = " ".join(imported_modules).lower()
    assert "posthog" not in joined
    assert "analytics" not in joined
    assert "events" not in joined
    # And no analytics emitter is bound on the module.
    assert not hasattr(tasks, "posthog_client")


# ===========================================================================
# Cross-cutting implementation contracts
# ===========================================================================
def test_row_lock_used_on_export_record_fetch(patch_session, captured_put):
    record = _make_record(status="Queued")
    session = patch_session(FakeSession(record=record, symbols=_make_symbols()))

    tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    assert ExportRecord in session.with_for_update_models


def test_db_session_closed_on_exception(patch_session, captured_put):
    record = _make_record(status="Queued")
    session = patch_session(FakeSession(record=record, symbols=_make_symbols()))
    captured_put.mock.side_effect = RuntimeError("S3 down")

    with pytest.raises(RuntimeError):
        tasks._process_export(EXPORT_ID, DRAWING_ID, USER_ID, "csv", allow_generating=False)

    assert session.closed is True


def test_task_has_acks_late_true():
    assert tasks.generate_export.acks_late is True
    assert tasks.generate_export.max_retries == 1
    assert tasks.generate_export.time_limit == 3600


def test_task_registered_on_export_queue():
    # Routed to the `export` queue via the `export.*` name prefix (§1.11).
    assert tasks.generate_export.name == "export.generate_export"
    assert tasks.generate_export.name in celery_app.tasks
    assert celery_app.conf.task_routes["export.*"]["queue"] == "export"


def test_publish_export_event_swallows_redis_errors(monkeypatch):
    """Fire-and-forget: a Redis publish failure is logged, never raised, so it
    cannot fail/retry a finished export (mirrors S1-D's publish_event)."""
    boom = MagicMock(name="redis")
    boom.publish.side_effect = ConnectionError("redis down")
    monkeypatch.setattr(tasks, "_redis_client", lambda: boom)

    # Must not raise.
    tasks.publish_export_event(f"export:status:{EXPORT_ID}", {"event": "export_complete"})
    boom.publish.assert_called_once()


def test_task_payload_kwarg_contract():
    """Load-bearing contract with S2-D's apply_async kwargs (must not drift)."""
    import inspect

    params = set(inspect.signature(tasks.generate_export.run).parameters)
    assert {"export_id", "drawing_id", "user_id", "format"} <= params


def test_user_id_sourced_from_payload_not_db(patch_session, captured_put, monkeypatch):
    # The DB record carries a DIFFERENT user_id; the published event must use
    # the payload value (§1.4 rule 5), proving no DB lookup for user_id.
    record = _make_record(status="Queued", user_id="ffffffff-9999-9999-9999-999999999999")
    patch_session(FakeSession(record=record, symbols=_make_symbols()))
    publish = MagicMock(name="publish_export_event")
    monkeypatch.setattr(tasks, "publish_export_event", publish)

    tasks._run_export(_make_self(retries=0), EXPORT_ID, DRAWING_ID, USER_ID, "csv")

    _, payload = publish.call_args.args
    assert payload["user_id"] == USER_ID
    assert payload["user_id"] != record.user_id
